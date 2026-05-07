# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt


def main() -> None:
    out_path = "Data_Dictionary_Table_3_ru.docx"
    doc = Document()
    doc.add_heading("Таблица 3. Словарь данных", level=1)

    entities = [
        ("Сущность user", [
            ("PK", "id", "INTEGER", "N", "Автозаполнение, первичный ключ"),
            ("", "username", "VARCHAR", "N", "Уникальный логин пользователя (email)"),
            ("", "password_hash", "VARCHAR", "N", "Хэш пароля"),
            ("", "role", "VARCHAR", "N", "Роль пользователя (administrator/manager/dentist)"),
            ("", "is_active", "BOOLEAN", "N", "Признак активности учетной записи"),
        ]),
        ("Сущность patient", [
            ("PK", "id", "INTEGER", "N", "Автозаполнение, первичный ключ"),
            ("", "full_name", "VARCHAR", "N", "ФИО пациента"),
            ("", "birth_date", "DATE", "Y", "Дата рождения"),
            ("", "phone", "VARCHAR", "Y", "Контактный телефон"),
            ("", "address", "VARCHAR", "Y", "Адрес проживания"),
            ("", "allergies", "TEXT", "Y", "Аллергии и противопоказания"),
            ("", "note", "TEXT", "Y", "Дополнительная заметка по пациенту"),
        ]),
        ("Сущность appointment", [
            ("PK", "id", "INTEGER", "N", "Идентификатор записи на прием"),
            ("FK", "patient_id", "INTEGER", "N", "ID пациента (patient.id)"),
            ("FK", "doctor_id", "INTEGER", "N", "ID врача (user.id)"),
            ("", "date", "DATE", "N", "Дата приема"),
            ("", "time", "TIME", "N", "Время приема"),
            ("", "status", "VARCHAR", "N", "Статус записи (active/cancelled/completed)"),
            ("", "comment", "TEXT", "Y", "Комментарий к записи"),
        ]),
        ("Сущность visit", [
            ("PK", "id", "INTEGER", "N", "Идентификатор визита"),
            ("FK", "patient_id", "INTEGER", "N", "ID пациента (patient.id)"),
            ("FK", "doctor_id", "INTEGER", "N", "ID врача (user.id)"),
            ("", "datetime", "TIMESTAMP", "N", "Дата и время визита"),
            ("", "complaints", "TEXT", "Y", "Жалобы пациента"),
            ("", "anamnesis", "TEXT", "Y", "Анамнез"),
            ("", "diagnosis", "TEXT", "Y", "Диагноз"),
            ("", "exam_results", "TEXT", "Y", "Результаты осмотра"),
            ("", "notes", "TEXT", "Y", "Дополнительные заметки"),
        ]),
        ("Сущность research", [
            ("PK", "id", "INTEGER", "N", "Идентификатор исследования"),
            ("FK", "patient_id", "INTEGER", "N", "ID пациента (patient.id)"),
            ("FK", "visit_id", "INTEGER", "Y", "ID визита (visit.id), если есть привязка"),
            ("", "datetime", "TIMESTAMP", "N", "Дата и время исследования"),
            ("", "type", "VARCHAR", "N", "Тип исследования (УЗИ, рентген, КТ и т.д.)"),
            ("", "result", "TEXT", "Y", "Результат исследования"),
        ]),
        ("Сущность prescription", [
            ("PK", "id", "INTEGER", "N", "Идентификатор назначения"),
            ("FK", "patient_id", "INTEGER", "N", "ID пациента (patient.id)"),
            ("FK", "visit_id", "INTEGER", "Y", "ID визита (visit.id), если есть привязка"),
            ("", "medication", "VARCHAR", "N", "Наименование препарата"),
            ("", "dosage", "VARCHAR", "Y", "Дозировка"),
            ("", "instructions", "TEXT", "Y", "Рекомендации по приему"),
            ("", "start_date", "DATE", "Y", "Дата начала приема"),
            ("", "end_date", "DATE", "Y", "Дата окончания приема"),
        ]),
        ("Сущность tooth", [
            ("PK", "id", "INTEGER", "N", "Идентификатор записи о зубе"),
            ("FK", "patient_id", "INTEGER", "N", "ID пациента (patient.id)"),
            ("", "number", "INTEGER", "N", "Номер зуба"),
            ("", "status", "VARCHAR", "Y", "Состояние зуба"),
            ("", "notes", "TEXT", "Y", "Примечания"),
        ]),
        ("Сущность servicecode", [
            ("PK", "id", "INTEGER", "N", "Идентификатор услуги"),
            ("", "code", "VARCHAR", "N", "Код услуги (МКБ-С-3)"),
            ("", "name", "VARCHAR", "N", "Название услуги"),
            ("", "description", "TEXT", "Y", "Описание услуги"),
        ]),
        ("Сущность work", [
            ("PK", "id", "INTEGER", "N", "Идентификатор наряда/работы"),
            ("FK", "visit_id", "INTEGER", "N", "ID визита (visit.id)"),
            ("FK", "service_code_id", "INTEGER", "Y", "ID услуги (servicecode.id)"),
            ("", "description", "TEXT", "Y", "Описание выполненных работ"),
            ("", "duration_minutes", "INTEGER", "Y", "Длительность в минутах"),
            ("", "materials", "TEXT", "Y", "Использованные материалы"),
            ("", "cost", "REAL", "Y", "Стоимость"),
            ("", "status", "VARCHAR", "N", "Статус работы (planned/in_progress/completed/cancelled)"),
            ("", "work_order_number", "VARCHAR", "Y", "Номер наряда"),
            ("", "tooth_numbers", "VARCHAR", "Y", "Номера зубов через запятую"),
            ("", "work_date", "TIMESTAMP", "N", "Дата и время выполнения работ"),
        ]),
        ("Сущность referral", [
            ("PK", "id", "INTEGER", "N", "Идентификатор направления"),
            ("FK", "patient_id", "INTEGER", "N", "ID пациента (patient.id)"),
            ("FK", "doctor_id", "INTEGER", "N", "ID врача (user.id)"),
            ("FK", "visit_id", "INTEGER", "Y", "ID визита (visit.id), если есть привязка"),
            ("", "referral_type", "VARCHAR", "N", "Тип направления (research/consultation)"),
            ("", "destination", "VARCHAR", "N", "Куда направлен пациент"),
            ("", "reason", "TEXT", "N", "Причина направления"),
            ("", "date", "TIMESTAMP", "N", "Дата и время направления"),
            ("", "status", "VARCHAR", "N", "Статус (created/completed/cancelled)"),
            ("", "result", "TEXT", "Y", "Результат выполнения направления"),
            ("", "notes", "TEXT", "Y", "Примечания"),
        ]),
    ]

    headers = ["Ключ", "Имя поля", "Тип данных", "Нулевые значения", "Дополнительное описание"]
    for title, rows in entities:
        doc.add_paragraph("")
        p = doc.add_paragraph(title)
        p.runs[0].bold = True
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
