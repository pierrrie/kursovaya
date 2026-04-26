from fastapi import APIRouter, Depends, HTTPException, status
from sqlmodel import Session, select
from typing import Dict, Any
from datetime import datetime
import json
from fastapi.responses import StreamingResponse, Response
from io import BytesIO
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.db.session import get_session
from app.auth.security import get_current_user
from app.models.user import User, UserRole
from app.models.patient import Patient
from app.models.visit import Visit
from app.models.work import Work
from app.models.research import Research
from app.models.prescription import Prescription
from app.models.tooth import Tooth
from app.models.referral import Referral

router = APIRouter()

@router.get("/patient-history/{patient_id}")
async def generate_patient_history(
    patient_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """Создание выписки из медицинской карты пациента"""
    if current_user.role not in [UserRole.DENTIST, UserRole.ADMIN]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Недостаточно прав"
        )
    
    # Получаем данные пациента
    patient = session.get(Patient, patient_id)
    if not patient:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Пациент не найден"
        )
    
    # Получаем все визиты пациента
    visits = session.exec(
        select(Visit)
        .where(Visit.patient_id == patient_id)
        .order_by(Visit.datetime.desc())
    ).all()
    
    # Получаем зубы пациента
    teeth = session.exec(
        select(Tooth)
        .where(Tooth.patient_id == patient_id)
        .order_by(Tooth.number)
    ).all()
    
    # Получаем исследования
    researches = session.exec(
        select(Research)
        .where(Research.patient_id == patient_id)
        .order_by(Research.datetime.desc())
    ).all()
    
    # Получаем назначения
    prescriptions = session.exec(
        select(Prescription)
        .where(Prescription.patient_id == patient_id)
        .order_by(Prescription.start_date.desc())
    ).all()
    
    # Получаем направления
    referrals = session.exec(
        select(Referral)
        .where(Referral.patient_id == patient_id)
        .order_by(Referral.date.desc())
    ).all()
    
    # Формируем зубную формулу
    tooth_formula = {}
    for tooth in teeth:
        tooth_formula[tooth.number] = {
            "status": tooth.status,
            "notes": tooth.notes
        }
    
    # Формируем отчет
    report = {
        "patient": {
            "id": patient.id,
            "full_name": patient.full_name,
            "birth_date": str(patient.birth_date) if patient.birth_date else None,
            "phone": patient.phone,
            "address": patient.address,
            "allergies": patient.allergies,
            "note": patient.note
        },
        "tooth_formula": tooth_formula,
        "visits": [
            {
                "id": visit.id,
                "datetime": str(visit.datetime),
                "doctor_id": visit.doctor_id,
                "complaints": visit.complaints,
                "anamnesis": visit.anamnesis,
                "diagnosis": visit.diagnosis,
                "exam_results": visit.exam_results,
                "notes": visit.notes
            } for visit in visits
        ],
        "researches": [
            {
                "id": research.id,
                "datetime": str(research.datetime),
                "type": research.type,
                "result": research.result
            } for research in researches
        ],
        "prescriptions": [
            {
                "id": prescription.id,
                "medication": prescription.medication,
                "dosage": prescription.dosage,
                "instructions": prescription.instructions,
                "start_date": str(prescription.start_date) if prescription.start_date else None,
                "end_date": str(prescription.end_date) if prescription.end_date else None
            } for prescription in prescriptions
        ],
        "referrals": [
            {
                "id": referral.id,
                "date": str(referral.date),
                "type": referral.referral_type,
                "destination": referral.destination,
                "reason": referral.reason,
                "status": referral.status,
                "result": referral.result
            } for referral in referrals
        ],
        "report_info": {
            "generated_at": datetime.now().isoformat(),
            "generated_by": current_user.username,
            "generated_by_role": current_user.role
        }
    }
    
    return report

@router.get("/patient-history/{patient_id}/word")
async def download_patient_history_word(
    patient_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """Скачать выписку из медицинской карты в формате Word"""
    if current_user.role not in [UserRole.DENTIST, UserRole.ADMIN]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Недостаточно прав"
        )
    
    # Получаем данные пациента
    patient = session.get(Patient, patient_id)
    if not patient:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Пациент не найден"
        )
    
    # Получаем все визиты пациента
    visits = session.exec(
        select(Visit)
        .where(Visit.patient_id == patient_id)
        .order_by(Visit.datetime.desc())
    ).all()
    
    # Создаем Word документ
    doc = Document()
    doc.add_heading('ВЫПИСКА ИЗ МЕДИЦИНСКОЙ КАРТЫ', 0)
    
    # Данные пациента
    doc.add_heading('Данные пациента', level=1)
    doc.add_paragraph(f'ФИО: {patient.full_name}')
    if patient.birth_date:
        doc.add_paragraph(f'Дата рождения: {patient.birth_date.strftime("%d.%m.%Y")}')
    doc.add_paragraph(f'Телефон: {patient.phone or "—"}')
    doc.add_paragraph(f'Адрес: {patient.address or "—"}')
    doc.add_paragraph(f'Аллергии: {patient.allergies or "—"}')
    
    # История визитов
    doc.add_heading('История визитов', level=1)
    for visit in visits:
        doc.add_heading(f'Визит от {visit.datetime.strftime("%d.%m.%Y %H:%M")}', level=2)
        doc.add_paragraph(f'Жалобы: {visit.complaints or "—"}')
        doc.add_paragraph(f'Анамнез: {visit.anamnesis or "—"}')
        doc.add_paragraph(f'Диагноз: {visit.diagnosis or "—"}')
        doc.add_paragraph(f'Результаты осмотра: {visit.exam_results or "—"}')
        doc.add_paragraph(f'Заметки: {visit.notes or "—"}')
        doc.add_paragraph('')  # Пустая строка
    
    # Сохраняем документ в память
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Возвращаем файл
    return Response(
        content=buffer.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )

@router.get("/medical-card/{patient_id}/word")
async def download_medical_card_word(
    patient_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """Скачать медицинскую карту стоматологического больного (форма 043/у) в формате Word"""
    if current_user.role not in [UserRole.DENTIST, UserRole.ADMIN]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Недостаточно прав"
        )
    
    # Получаем данные пациента
    patient = session.get(Patient, patient_id)
    if not patient:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Пациент не найден"
        )
    
    # Получаем все визиты пациента
    visits = session.exec(
        select(Visit)
        .where(Visit.patient_id == patient_id)
        .order_by(Visit.datetime.desc())
    ).all()
    
    # Получаем зубы пациента
    teeth = session.exec(
        select(Tooth)
        .where(Tooth.patient_id == patient_id)
        .order_by(Tooth.number)
    ).all()
    
    # Создаем Word документ
    doc = Document()
    doc.add_heading('МЕДИЦИНСКАЯ КАРТА СТОМАТОЛОГИЧЕСКОГО БОЛЬНОГО', 0)
    doc.add_heading('Форма 043/у', 1)
    
    # Данные пациента
    doc.add_heading('Личные данные пациента', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Поле'
    hdr_cells[1].text = 'Значение'
    
    rows_data = [
        ('ФИО', patient.full_name),
        ('Дата рождения', patient.birth_date.strftime("%d.%m.%Y") if patient.birth_date else "—"),
        ('Телефон', patient.phone or "—"),
        ('Адрес', patient.address or "—"),
        ('Аллергии', patient.allergies or "—"),
        ('Примечания', patient.note or "—"),
    ]
    
    for field, value in rows_data:
        row_cells = table.add_row().cells
        row_cells[0].text = field
        row_cells[1].text = value
    
    # Зубная формула
    doc.add_heading('Зубная формула', level=1)
    if teeth:
        tooth_table = doc.add_table(rows=1, cols=4)
        tooth_table.style = 'Table Grid'
        
        hdr_cells = tooth_table.rows[0].cells
        hdr_cells[0].text = 'Номер зуба'
        hdr_cells[1].text = 'Статус'
        hdr_cells[2].text = 'Примечания'
        hdr_cells[3].text = 'Дата последнего обновления'
        
        for tooth in teeth:
            row_cells = tooth_table.add_row().cells
            row_cells[0].text = str(tooth.number)
            row_cells[1].text = tooth.status or "—"
            row_cells[2].text = tooth.notes or "—"
            updated_at = getattr(tooth, "updated_at", None)
            row_cells[3].text = updated_at.strftime("%d.%m.%Y") if updated_at else "—"
    else:
        doc.add_paragraph('Зубная формула не заполнена')
    
    # История лечения
    doc.add_heading('История лечения', level=1)
    for visit in visits:
        doc.add_heading(f'Визит от {visit.datetime.strftime("%d.%m.%Y %H:%M")}', level=2)
        
        visit_table = doc.add_table(rows=1, cols=2)
        visit_table.style = 'Table Grid'
        
        hdr_cells = visit_table.rows[0].cells
        hdr_cells[0].text = 'Параметр'
        hdr_cells[1].text = 'Значение'
        
        visit_data = [
            ('Жалобы', visit.complaints or "—"),
            ('Анамнез', visit.anamnesis or "—"),
            ('Диагноз', visit.diagnosis or "—"),
            ('Результаты осмотра', visit.exam_results or "—"),
            ('Заметки', visit.notes or "—"),
        ]
        
        for param, value in visit_data:
            row_cells = visit_table.add_row().cells
            row_cells[0].text = param
            row_cells[1].text = value
        
        doc.add_paragraph('')  # Пустая строка между визитами
    
    # Сохраняем документ в память
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Возвращаем файл
    return Response(
        content=buffer.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )

@router.get("/visit-report/{visit_id}")
async def generate_visit_report(
    visit_id: int,
    session: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """Создание отчета по визиту (форма 043/у)"""
    if current_user.role not in [UserRole.DENTIST, UserRole.ADMIN]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Недостаточно прав"
        )
    
    # Получаем визит
    visit = session.get(Visit, visit_id)
    if not visit:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Визит не найден"
        )
    
    # Получаем пациента
    patient = session.get(Patient, visit.patient_id)
    
    # Получаем врача
    from app.models.user import User as Doctor
    doctor = session.get(Doctor, visit.doctor_id)
    
    # Получаем работы по визиту
    works = session.exec(
        select(Work)
        .where(Work.visit_id == visit_id)
    ).all()
    
    # Формируем наряд на работу
    work_order = {
        "works": [
            {
                "description": work.description,
                "duration_minutes": work.duration_minutes,
                "materials": work.materials,
                "cost": work.cost,
                "tooth_numbers": work.tooth_numbers,
                "work_order_number": work.work_order_number
            } for work in works
        ],
        "total_cost": sum(work.cost for work in works if work.cost),
        "total_duration": sum(work.duration_minutes for work in works if work.duration_minutes)
    }
    
    # Формируем отчет по форме 043/у
    report = {
        "form_number": "043/у",
        "form_name": "Медицинская карта стоматологического больного",
        "patient": {
            "full_name": patient.full_name if patient else "Неизвестно",
            "birth_date": str(patient.birth_date) if patient and patient.birth_date else None,
            "address": patient.address if patient else None,
            "phone": patient.phone if patient else None
        },
        "visit": {
            "datetime": str(visit.datetime),
            "doctor": doctor.username if doctor else "Неизвестно",
            "complaints": visit.complaints,
            "anamnesis": visit.anamnesis,
            "diagnosis": visit.diagnosis,
            "exam_results": visit.exam_results,
            "notes": visit.notes
        },
        "work_order": work_order,
        "report_info": {
            "generated_at": datetime.now().isoformat(),
            "generated_by": current_user.username
        }
    }
    
    return report